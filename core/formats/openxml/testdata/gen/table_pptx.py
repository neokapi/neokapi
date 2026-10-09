# /// script
# requires-python = ">=3.11"
# dependencies = ["python-pptx>=1.0,<2"]
# ///
"""Write testdata/table.pptx: a two-slide deck whose second slide holds a table.

The generator sits in testdata/gen/ because the validity tests open every file
directly under testdata/ as a package.

The first slide is a title slide. The second has the title "Fees at a glance", a
4 × 3 table whose last row merges its first two cells, and speaker notes. The
reader tests hold the slide table to a table Group of row Groups with cell
roles, and the deck to one read per slide with no layout prompt text.

    uv run core/formats/openxml/testdata/gen/table_pptx.py
"""

import datetime as dt
import io
import pathlib
import zipfile

from pptx import Presentation
from pptx.util import Inches

HERE = pathlib.Path(__file__).resolve().parent
STAMP = dt.datetime(2026, 9, 1, 9, 0, 0)


def normalise_zip(data: bytes) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as dst:
        for item in src.infolist():
            info = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = item.external_attr
            dst.writestr(info, src.read(item.filename))
    return out.getvalue()


def main() -> None:
    prs = Presentation()
    cp = prs.core_properties
    cp.author = "neokapi"
    cp.last_modified_by = "neokapi"
    cp.title = "Table deck"
    cp.created = STAMP
    cp.modified = STAMP
    cp.revision = 1

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "Table deck"
    s.placeholders[1].text = "A slide table for the reader tests"

    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Fees at a glance"
    rows = [
        ("Plan", "Monthly", "Payout"),
        ("Starter", "€0", "Monthly"),
        ("Growth", "€49", "Weekly"),
        ("Enterprise, custom terms", "", "Daily"),
    ]
    shape = s.shapes.add_table(len(rows), 3, Inches(0.8), Inches(1.8), Inches(8.4), Inches(2.4))
    table = shape.table
    for r, cells in enumerate(rows):
        for c, text in enumerate(cells):
            table.cell(r, c).text = text
    table.cell(3, 0).merge(table.cell(3, 1))
    s.notes_slide.notes_text_frame.text = "Keep this slide short."

    buf = io.BytesIO()
    prs.save(buf)
    out = HERE.parent / "table.pptx"
    out.write_bytes(normalise_zip(buf.getvalue()))
    print("wrote", out)


if __name__ == "__main__":
    main()
